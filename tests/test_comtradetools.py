"""Tests for comtradetools.py, mirroring the calls made in the notebooks.

Notebook patterns covered:
- 0-comtrade-setup-first.ipynb: setup / get_api_key / init
- cn|hk|mo|tw_plp_import_export.ipynb: getFinalData with notebook kwargs,
  year_range, split_period
- country_trade_profile.ipynb: get_trade_flows, total_rank_perc, make_format
- cn_plp_commodities.ipynb: excel_col_autowidth / excel_format_currency /
  excel_format_percent, encode_country
- comtrade-api.ipynb, isaggregate_bug.ipynb: decode_country, checkAggregateValues

No test performs network access: the API layer is replaced by fakes via
monkeypatch, and caching is redirected to pytest tmp dirs.
"""

import pandas as pd
import pytest

import comtradetools


def make_api_df(periods=("2020",), reporter=156, partners=(24, 76), flow="M", value=100.0):
    """Small DataFrame mimicking a comtradeapicall.getFinalData result."""
    rows = [
        {
            "period": per,
            "refYear": int(per),
            "reporterCode": reporter,
            "partnerCode": p,
            "partner2Code": 0,
            "cmdCode": "TOTAL",
            "flowCode": flow,
            "primaryValue": value,
        }
        for per in periods
        for p in partners
    ]
    return pd.DataFrame(rows)


class FakeAPI:
    """Records calls and returns a preset DataFrame (None simulates a failed call)."""

    def __init__(self):
        self.calls = []
        self.next_result = None

    def __call__(self, *p, **kwp):
        self.calls.append(kwp.copy())
        if self.next_result is None:
            return None
        return self.next_result.copy()


@pytest.fixture
def fake_api(monkeypatch, tmp_path):
    """Replace the rate-limited inner wrapper with a recorder, and redirect the cache."""
    fake = FakeAPI()
    fake.next_result = make_api_df()
    monkeypatch.setattr(comtradetools, "comtradeapicall_getFinalData", fake)
    monkeypatch.setattr(comtradetools, "CACHE_DIR", str(tmp_path))
    monkeypatch.setattr(comtradetools, "MAX_SLEEP", 0)  # keep retry tests fast
    return fake


# ---------------------------------------------------------------------------
# setup / get_api_key / init — 0-comtrade-setup-first.ipynb
# ---------------------------------------------------------------------------


class TestSetupInit:
    def test_setup_creates_dirs_and_config(self, tmp_path):
        comtradetools.setup(
            support_dir=str(tmp_path / "support"),
            cache_dir=str(tmp_path / "cache"),
            config_file=str(tmp_path / "config.ini"),
        )
        assert (tmp_path / "support").is_dir()
        assert (tmp_path / "cache").is_dir()
        assert (tmp_path / "config.ini").is_file()

    def test_get_api_key_reads_config(self, ctt):
        assert ctt.get_api_key() == "TESTKEY"

    def test_init_populates_reference_globals(self, ctt):
        assert ctt.INIT_DONE is True
        assert ctt.COUNTRY_CODES[156] == "China"
        assert ctt.COUNTRY_CODES_REVERSE["China"] == 156
        assert len(ctt.HS_CODES) > 0
        assert len(ctt.HS_CODES_L2) > 0
        assert set(ctt.PLP_CODES) == set(ctt.m49_plp)
        assert "M" in ctt.FLOWS_CODES

    def test_init_is_noop_unless_forced(self, ctt):
        sentinel = {"sentinel": 1}
        original = ctt.COUNTRY_CODES
        ctt.COUNTRY_CODES = sentinel
        try:
            ctt.init("TESTKEY")  # must return early, keeping the sentinel
            assert ctt.COUNTRY_CODES is sentinel
        finally:
            ctt.COUNTRY_CODES = original

    def test_get_url_selects_endpoint(self, ctt):
        assert ctt.get_url("SOMEKEY") == ctt.BASE_URL_API
        assert ctt.get_url(None) == ctt.BASE_URL_PREVIEW
        assert ctt.get_url("APIKEYHERE") == ctt.BASE_URL_PREVIEW


# ---------------------------------------------------------------------------
# encode/decode_country — cn_plp_commodities.ipynb, comtrade-api.ipynb
# ---------------------------------------------------------------------------


class TestCountryCodecs:
    def test_encode_country(self, ctt):
        assert ctt.encode_country("China") == 156
        assert ctt.encode_country("Angola") == 24

    def test_decode_country(self, ctt):
        assert ctt.decode_country(156) == "China"
        assert ctt.decode_country(ctt.m49_brazil) == "Brazil"

    def test_unknown_passes_through(self, ctt):
        assert ctt.encode_country("Atlantis") == "Atlantis"
        assert ctt.decode_country(9999) == 9999


# ---------------------------------------------------------------------------
# period helpers — all import/export notebooks
# ---------------------------------------------------------------------------


class TestPeriodHelpers:
    def test_year_range(self):
        assert comtradetools.year_range(2020, 2022) == "2020,2021,2022"
        assert comtradetools.year_range(2020, 2020) == "2020"

    def test_split_period_chunks_at_12(self):
        period = comtradetools.year_range(2000, 2012)  # 13 years
        chunks = comtradetools.split_period(period)
        assert len(chunks) == 2
        assert chunks[0] == comtradetools.year_range(2000, 2011)
        assert chunks[1] == "2012"

    def test_split_period_short(self):
        assert comtradetools.split_period("2020,2021") == ["2020,2021"]

    def test_get_year_intervals(self):
        assert comtradetools.get_year_intervals([2018, 2019, 2021]) == [
            "2018-2019",
            "2021-2021",
        ]
        assert comtradetools.get_year_intervals([2020]) == ["2020-2020"]


# ---------------------------------------------------------------------------
# getFinalData — cn|hk|mo|tw_plp_import_export.ipynb, cn_plp_commodities.ipynb
# ---------------------------------------------------------------------------


class TestGetFinalData:
    # kwargs as in cn_plp_import_export.ipynb
    NOTEBOOK_KWARGS = dict(
        typeCode="C",
        freqCode="A",
        clCode="HS",
        cmdCode="TOTAL",
        flowCode="M",
        reporterCode=156,
        partner2Code=0,
        customsCode="C00",
        motCode="0",
        includeDesc=True,
    )

    def test_notebook_style_call(self, fake_api):
        fake_api.next_result = make_api_df(periods=("2020", "2021"))
        df = comtradetools.getFinalData(
            "TESTKEY",
            partnerCode=comtradetools.m49_plp_list,
            period=comtradetools.year_range(2020, 2021),
            **self.NOTEBOOK_KWARGS,
        )
        assert len(fake_api.calls) == 1
        assert len(df) == 2 * 2  # 2 periods x 2 partners in the fake
        assert df["primaryValue"].sum() > 0

    def test_periods_split_into_chunks(self, fake_api):
        comtradetools.getFinalData(
            "TESTKEY",
            partnerCode=76,
            period=comtradetools.year_range(2000, 2012),  # 13 years > 12 per call
            **self.NOTEBOOK_KWARGS,
        )
        assert len(fake_api.calls) == 2
        assert fake_api.calls[0]["period"] == comtradetools.year_range(2000, 2011)
        assert fake_api.calls[1]["period"] == "2012"

    def test_partner2code_defaults_to_zero(self, fake_api):
        kwargs = dict(self.NOTEBOOK_KWARGS)
        del kwargs["partner2Code"]
        comtradetools.getFinalData("TESTKEY", partnerCode=76, period="2020", **kwargs)
        assert fake_api.calls[0]["partner2Code"] == 0

    def test_period_is_required(self, fake_api):
        with pytest.raises(ValueError, match="Period is required"):
            comtradetools.getFinalData("TESTKEY", partnerCode=76, **self.NOTEBOOK_KWARGS)

    def test_second_identical_call_uses_cache(self, fake_api):
        kwargs = dict(self.NOTEBOOK_KWARGS, partnerCode=76, period="2020")
        first = comtradetools.getFinalData("TESTKEY", **kwargs)
        second = comtradetools.getFinalData("TESTKEY", **kwargs)
        assert len(fake_api.calls) == 1  # second call served from cache
        pd.testing.assert_frame_equal(first, second)

    def test_cache_false_forces_refetch(self, fake_api):
        kwargs = dict(self.NOTEBOOK_KWARGS, partnerCode=76, period="2020")
        comtradetools.getFinalData("TESTKEY", **kwargs)
        comtradetools.getFinalData("TESTKEY", cache=False, **kwargs)
        assert len(fake_api.calls) == 2

    def test_csv_list_splitting(self, fake_api):
        """Oversized CSV lists are batched (API ~2000-char URL limit); each
        batch has its own cache entries and results are concatenated."""
        codes = ",".join(f"{100000 + i}" for i in range(178))
        # period_size=1 as in country_trade_profile.ipynb §2.5/§3.5
        kwargs = dict(self.NOTEBOOK_KWARGS, partnerCode=None,
                      period="2020,2021", period_size=1, cmdCode=codes)
        df = comtradetools.getFinalData("TESTKEY", **kwargs)
        # 2 batches (100 + 78) x 2 yearly chunks = 4 inner calls
        assert len(fake_api.calls) == 4
        sizes = sorted(len(c["cmdCode"].split(",")) for c in fake_api.calls)
        assert sizes == [78, 78, 100, 100]
        assert len(df) == 8  # fake returns 2 rows per call
        comtradetools.getFinalData("TESTKEY", **kwargs)
        assert len(fake_api.calls) == 4  # repeat call fully served from cache

    def test_remove_world_drops_partnercode_zero(self, fake_api, monkeypatch):
        fake_api.next_result = make_api_df(partners=(0, 76, 24))
        df = comtradetools.getFinalData(
            "TESTKEY",
            **self.NOTEBOOK_KWARGS,
            partnerCode=None,
            period="2020",
            remove_world=True,
        )
        assert set(df["partnerCode"]) == {76, 24}

    def test_empty_result_returns_empty_dataframe(self, fake_api):
        fake_api.next_result = pd.DataFrame()
        df = comtradetools.getFinalData(
            "TESTKEY", **self.NOTEBOOK_KWARGS, partnerCode=76, period="2020"
        )
        assert isinstance(df, pd.DataFrame)
        assert df.empty

    def test_none_result_retries_then_raises(self, fake_api):
        fake_api.next_result = None
        with pytest.raises(IOError):
            comtradetools.getFinalData(
                "TESTKEY", **self.NOTEBOOK_KWARGS, partnerCode=76, period="2020"
            )
        # 1 initial + MAX_RETRIES retries of the "empty result" loop
        assert len(fake_api.calls) == 1 + comtradetools.MAX_RETRIES


# ---------------------------------------------------------------------------
# get_trade_flows — country_trade_profile.ipynb
# ---------------------------------------------------------------------------


def trade_flows_fake(monkeypatch):
    """Fake getFinalData dispatching on reporter/partner/flow like the API would."""
    calls = []

    def fake(*p, **kwp):
        calls.append(kwp.copy())
        periods = tuple(kwp["period"].split(","))
        flow = kwp["flowCode"]
        if kwp["reporterCode"] == 24:  # country of interest reports
            value = 100.0 if flow == "M" else 80.0
            if kwp.get("partnerCode") is None:
                # partnerCode=None response: World row (partnerCode=0) with the
                # total alongside individual partner rows — get_trade_flows
                # filters to the World row when partners == 0.
                world = make_api_df(periods, reporter=24, partners=(0,),
                                    flow=flow, value=value)
                rest = make_api_df(periods, reporter=24, partners=(76,),
                                   flow=flow, value=value / 2)
                return pd.concat([world, rest], ignore_index=True)
            return make_api_df(periods, reporter=24, partners=(76,), flow=flow,
                               value=value)
        # partners report (mirror calls): partnerCode is the country of interest
        value = 90.0 if flow == "M" else 70.0
        return make_api_df(periods, reporter=76, partners=(24,), flow=flow, value=value)

    monkeypatch.setattr(comtradetools, "getFinalData", fake)
    return calls


class TestGetTradeFlows:
    # call pattern as in country_trade_profile.ipynb
    def test_symmetric_flows_and_balance(self, monkeypatch):
        calls = trade_flows_fake(monkeypatch)
        tb = comtradetools.get_trade_flows(
            countryOfInterest=24,
            period="2020,2021",
            partners="76,132",
            period_size=1,
            retry_if_empty=False,
            symmetric_values=True,
        )
        assert len(calls) == 4  # M, X, and the two mirror calls
        row = tb.loc["2020"]
        assert row["M"] == 100.0
        assert row["X"] == 80.0
        assert row["X<M"] == 90.0
        assert row["M<X"] == 70.0
        assert row["trade_balance (X-M)"] == pytest.approx(-20.0)
        assert row["trade_balance (X<M-M)"] == pytest.approx(-10.0)
        assert row["trade_volume (X+M)"] == pytest.approx(180.0)
        assert row["trade_volume (X<M+M<X)"] == pytest.approx(160.0)

    def test_no_symmetric_values(self, monkeypatch):
        calls = trade_flows_fake(monkeypatch)
        tb = comtradetools.get_trade_flows(
            countryOfInterest=24, period="2020", symmetric_values=False
        )
        assert len(calls) == 2
        assert "X" in tb.columns and "M" in tb.columns
        assert "X<M" not in tb.columns

    def test_missing_exports_skips_derived_columns(self, monkeypatch):
        trade_flows_fake(monkeypatch)

        original = comtradetools.getFinalData

        def fake_no_exports(*p, **kwp):
            if kwp["reporterCode"] == 24 and kwp["flowCode"] == "X":
                return pd.DataFrame()
            return original(*p, **kwp)

        monkeypatch.setattr(comtradetools, "getFinalData", fake_no_exports)
        tb = comtradetools.get_trade_flows(
            countryOfInterest=24, period="2020", symmetric_values=True
        )
        assert "X" not in tb.columns
        assert "trade_balance (X-M)" not in tb.columns  # guarded
        assert "trade_balance (X<M-M)" in tb.columns


# ---------------------------------------------------------------------------
# DataFrame utilities — country_trade_profile.ipynb
# ---------------------------------------------------------------------------


class TestDataFrameUtils:
    SAMPLE = pd.DataFrame(
        {
            "year": [2020, 2020, 2020, 2021, 2021, 2021],
            "flow": ["M"] * 6,
            "partner": ["A", "B", "C", "A", "B", "C"],
            "primaryValue": [50.0, 30.0, 20.0, 40.0, 40.0, 20.0],
        }
    )

    def test_subtotal(self):
        out = comtradetools.subtotal(self.SAMPLE, ["year"], "primaryValue")
        assert out.tolist() == [100.0] * 3 + [100.0] * 3

    def test_rank_dense_descending(self):
        out = comtradetools.rank(self.SAMPLE, ["year"], "primaryValue")
        assert out.tolist() == [1, 2, 3, 1, 1, 2]

    def test_total_rank_perc(self):
        out = comtradetools.total_rank_perc(
            self.SAMPLE.copy(),
            groupby=["year", "flow", "partner"],
            col="primaryValue",
            prefix="partner",
        )
        row = out[(out.year == 2020) & (out.partner == "A")].iloc[0]
        assert row["partner_sum"] == 50.0
        assert row["partner_rank"] == 1
        assert row["partner_perc"] == pytest.approx(0.5)
        assert row["partner_upper_sum"] == 100.0
        assert row["partner_upper_perc"] == pytest.approx(0.5)
        # dense rank tie in 2021
        ranks_2021 = out[out.year == 2021]["partner_rank"].tolist()
        assert ranks_2021 == [1, 1, 2]

    def test_make_format(self):
        fmt = comtradetools.make_format(["a_perc", "b_sum", "primaryValue", "other"])
        assert fmt["a_perc"] == "{0:.3%}"
        assert fmt["b_sum"] == "${0:,.0f}"
        assert fmt["primaryValue"] == "${0:,.0f}"
        assert "other" not in fmt


# ---------------------------------------------------------------------------
# Excel helpers — cn_plp_commodities.ipynb
# ---------------------------------------------------------------------------


class TestExcelHelpers:
    DF = pd.DataFrame(
        {
            "country": ["Angola", "Brazil"],
            "primaryValue": [1234567.0, 8910111.0],
            "perc_x": [0.25, 0.75],
        }
    )

    def test_excel_roundtrip(self, tmp_path):
        from openpyxl import load_workbook

        path = tmp_path / "out.xlsx"
        df = self.DF.set_index("country")
        with pd.ExcelWriter(path, engine="xlsxwriter") as writer:
            df.to_excel(writer, sheet_name="data")
            comtradetools.excel_col_autowidth(df, writer, sheet="data")
            comtradetools.excel_format_currency(
                df, writer, sheet="data", columns=["primaryValue"]
            )
            comtradetools.excel_format_percent(df, writer, sheet="data", columns=["perc_x"])

        ws = load_workbook(path)["data"]
        # xlsxwriter stores a slightly adjusted width; assert it matches within tolerance
        assert ws.column_dimensions["A"].width == pytest.approx(len("country"), abs=1.0)
        assert "$" in ws["B2"].number_format
        assert "%" in ws["C2"].number_format


# ---------------------------------------------------------------------------
# checkAggregateValues — isaggregate_bug.ipynb
# ---------------------------------------------------------------------------


class TestCheckAggregateValues:
    def test_flags_parent_codes(self):
        df = pd.DataFrame({"cmdCode": ["01", "0101", "02"], "v": [10, 10, 5]})
        out = comtradetools.checkAggregateValues(df.copy(), "cmdCode")
        flags = out["isCmdAggregate"]
        assert flags.iloc[0] == True  # "01" is parent of "0101"  # noqa: E712
        assert flags.iloc[1] == False  # noqa: E712
        assert pd.isna(flags.iloc[2])  # last row is never assigned (documented behavior)

    def test_duplicate_codes_warn(self):
        df = pd.DataFrame({"cmdCode": ["0101", "0101"], "v": [1, 1]})
        with pytest.warns(UserWarning, match="duplicated"):
            comtradetools.checkAggregateValues(df.copy(), "cmdCode")
